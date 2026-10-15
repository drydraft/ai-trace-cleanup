#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""ai-trace-cleanup · 自检（回归闸门）

    python tests/selftest.py

它做的事，就是「这套规则还能不能用」的全部证明：
  T1 脏样本能生成，且 50 条规则**条条命中**（加了规则却没加样本 → 直接失败）
  T2 scan → apply → guard 全链路退出码为 0
  T3 关键替换逐条比对（跨 run 删空格这种核心能力，必须逐字对）
  T4 清理后文档能被 python-docx 独立打开（第三方解析器交叉验证）
  T5 清理后机械层零残留（幂等）
  T6 保护区逐字节未变、非文本部件 CRC 未变
  T7 二次 apply 无事可做（真幂等，不只是"扫不到"）
"""

from __future__ import annotations

import io
import json
import os
import subprocess
import sys
import zipfile

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
sys.path.insert(0, os.path.join(ROOT, "scripts"))

from docx_text import read_docx, iter_paragraphs, u  # noqa: E402
from rules_core import RULES  # noqa: E402
from rules_report import (CAT_RULE, C_PATTERNS, L_STRUCT, M_PATTERNS,  # noqa: E402
                          load_lexicon)

PY = sys.executable
SAMPLE = os.path.join(HERE, "sample-dirty.docx")
CLEAN = os.path.join(HERE, "clean.docx")


def run(*args) -> tuple:
    p = subprocess.run([PY] + list(args), capture_output=True, text=True,
                       encoding="utf-8", errors="replace", cwd=ROOT)
    return p.returncode, (p.stdout or "") + (p.stderr or "")


def rule_spec() -> list:
    from rules_report import REPORT_RULE_META
    lcats = [CAT_RULE[c][0] for c in load_lexicon() if c in CAT_RULE]
    return sorted(set(
        [r[0] for r in RULES] + [m[0] for m in M_PATTERNS]
        + [s[0] for s in L_STRUCT] + list(REPORT_RULE_META)
        + lcats + [c[0] for c in C_PATTERNS]))


def paras(path):
    doc = read_docx(path)
    out = {}
    for part in doc.text_parts:
        for p in iter_paragraphs(doc.xml(part), part)[0]:
            out[(part, p.text.strip())] = p
    return doc, out


#: 关键替换的期望结果（改规则时这里必须跟着改，改不了说明规则坏了）
EXPECT = {
    "本次巡检共 发现 3 个问题，已闭环 2 项。": "本次巡检共发现3个问题，已闭环2项。",
    "2026 年第三季度巡检工作总结": "2026年第三季度巡检工作总结",
    "检查完成, 未发现异常.": "检查完成，未发现异常。",
    "省略号... 还有。。。两种写法": "省略号……还有……两种写法",
    "发现了严重问题！！为什么呢？？": "发现了严重问题！为什么呢？",
    "占比 50 ％，全角数字１２３": "占比50％，全角数字123",
    "括号内空格（ 说明文字 ）与尾随空格 ": "括号内空格（说明文字）与尾随空格",
    "共 3 个问题": "共3个问题",
    "已处理 2 个,剩余 1 个": "已处理2个，剩余1个",
    "文本框里也有 3 个空格问题": "文本框里也有3个空格问题",
    "巡检工作总结 共 5 页": "巡检工作总结共5页",
    # 文本级保护区：路径 / 命令 / 版本 / 时间 / 单位 / 代码段原样保留，
    # 而同段里第 3 行的「共 3 处」照样要修 —— 两者同时成立才叫保护正确。
    "保护路径 .github/workflows/ 共 3 个文件。": "保护路径 .github/workflows/ 共3个文件。",
    r"路径 C:\Users\Administrator\.pi\agent\skills\ 下共 2 个目录。":
        r"路径 C:\Users\Administrator\.pi\agent\skills\ 下共2个目录。",
    "命令 python scripts/docx_scan.py --all 与版本 v0.1.0 都不许改，但共 3 处空格要删。":
        "命令 python scripts/docx_scan.py --all 与版本 v0.1.0都不许改，但共3处空格要删。",
    # 这一句最能说明取舍：`单位 10` 的交界空格要删（S01），而 `10 kg` 里的空格要留（掩码保护）。
    "时间 9:30、单位 10 kg、邮箱 a@b.com、代码 `x = 1` 都不能动。":
        "时间9:30、单位10 kg、邮箱 a@b.com、代码 `x = 1` 都不能动。",
}

#: T8 文本级保护区：这些 token 必须**整体**落在掩码里（一个字符都不许漏）
MUST_MASK = [
    ("保护 .github/workflows/ 这个路径", ".github/workflows"),
    (r"放在C:\Users\Administrator\.pi\agent\skills\ 下",
     "C:\\Users\\Administrator\\.pi\\agent\\skills\\"),
    ("见 ~/.pi/agent/AGENTS.md 文件", "~/.pi/agent/AGENTS.md"),
    ("运行 python scripts/docx_scan.py --all 即可", "--all"),
    ("全量跑 docx_scan.py --rules P01,P03 即可", "--rules"),
    ("用 `code span` 与 run-all.ps1 测", "`code span`"),
    ("用 `code span` 与 run-all.ps1 测", "run-all.ps1"),
    ("参考 docs/MAINTENANCE.md 与 tests/selftest.py", "docs/MAINTENANCE.md"),
    ("参考 docs/MAINTENANCE.md 与 tests/selftest.py", "tests/selftest.py"),
    ("网站 https://github.com/drydraft/ai-trace-cleanup 上",
     "https://github.com/drydraft/ai-trace-cleanup"),
    ("邮箱 a@b.com 可用", "a@b.com"),
    ("函数 os.path.join 与 getElementById", "os.path.join"),
    ("函数 os.path.join 与 getElementById", "getElementById"),
    ("时间 9:30 与日期 2026-10-13", "9:30"),
    ("时间 9:30 与日期 2026-10-13", "2026-10-13"),
    ("版本 v0.1.0 与小数 1.5", "v0.1.0"),
    ("版本 v0.1.0 与小数 1.5", "1.5"),
    ("单位 10 kg 与 100 MB", "10 kg"),
    ("单位 10 kg 与 100 MB", "100 MB"),
    ("变量 %USERPROFILE% 与 $HOME", "%USERPROFILE%"),
    ("数据 data.7z 与 ci.yml", "data.7z"),
    ("斜杠 and/or 与 TCP/IP", "and/or"),
    ("key: value 这一行", "key:"),
    ("Note: 以下说明", "Note:"),
    ("编号 JIRA-1234 与 UTF-8", "JIRA-1234"),
    ("实体 &nbsp; 与 &#160;", "&nbsp;"),
]

#: T8 反向：这些必须**不**被保护，否则就把该做的修复挡住了
MUST_NOT_MASK = [
    "检查完成, 未发现异常.",              # P01 要把 , 和 . 转全角
    "共 3 个问题，已处理 2 个",            # S01 要删中数之间那两个空格
    "占比 50 ％，全角数字１２３",            # P07 要转全角数字，S03 要删空格
    "省略号... 还有。。。两种写法",         # P03 要统一成 ……
    "数量 10-20 台，另一处写作 A--B",       # P04 要把 -- 改破折号
    "时间 9：30 使用了全角冒号",            # P08 要把全角冒号改半角
    "温度应在 3～5 度之间。",              # P09 要把 ～ 改 -
    "百分比 50 %",                        # S06 要删 % 前的空格
    "号 2. UPS 电池组告警",                # S02 要报「UPS 电池组」之间那个空格
    "作为一个 AI 语言模型",                # S02 与 C01 都要报这里
]


def main() -> int:
    u()
    fails, notes = [], []

    def check(ok, label, detail=""):
        print("%s %s%s" % ("✅" if ok else "❌", label, ("　" + detail) if detail else ""))
        if not ok:
            fails.append(label + (("　" + detail) if detail else ""))
        return ok

    print("── T1 生成脏样本并检查规则覆盖率 ──")
    rc, out = run(os.path.join(HERE, "make_sample.py"), SAMPLE)
    check(rc == 0, "脏样本生成", out.strip().splitlines()[-1] if out.strip() else "")
    rc, out = run(os.path.join(ROOT, "scripts", "docx_scan.py"), SAMPLE, "--all",
                  "--out", os.path.join(HERE, "scan-all"), "--quiet")
    check(rc == 3, "审计能报出硬阻（退出码 3）", "实际 %d" % rc)
    res = json.load(open(os.path.join(HERE, "scan-all", "report.json"), encoding="utf-8"))
    hit = {i["rule"] for i in res["issues"]}
    spec = rule_spec()
    uncovered = [r for r in spec if r not in hit]
    check(not uncovered, "每条规则都有样本命中（%d 条）" % len(spec),
          "未命中：%s" % uncovered)

    print("\n── T2 全链路：scan → apply → guard ──")
    for stale in (CLEAN, CLEAN + ".applied.json"):
        if os.path.isfile(stale):
            os.remove(stale)
    rc, out = run(os.path.join(ROOT, "scripts", "docx_apply.py"), SAMPLE,
                  "--out", CLEAN, "--apply")
    check(rc == 0, "apply 落盘成功（校验通过才写）", out.strip().splitlines()[0] if out.strip() else "")
    rc, out = run(os.path.join(ROOT, "scripts", "docx_guard.py"), SAMPLE, CLEAN,
                  "--applied", CLEAN + ".applied.json")
    check(rc == 0, "guard 七项校验全过")
    for line in out.splitlines():
        if line.strip().startswith(("✅", "❌", "⏭")):
            notes.append("     " + line.strip())

    print("\n── T3 关键替换逐条比对 ──")
    _, after = paras(CLEAN)
    for before, want in EXPECT.items():
        got = next((k[1] for k in after if k[1] == want), None)
        check(got is not None, "「%s」→「%s」" % (before[:18], want[:18]))

    print("\n── T4 第三方解析器交叉验证 ──")
    try:
        import docx as _docx
        d = _docx.Document(CLEAN)
        check(len(d.paragraphs) > 50 and len(d.tables) == 1,
              "python-docx 能独立打开", "段落 %d / 表格 %d" % (len(d.paragraphs), len(d.tables)))
    except ImportError:
        print("⏭️ 未装 python-docx，跳过（不影响结论）")

    print("\n── T5 幂等 ──")
    rc, out = run(os.path.join(ROOT, "scripts", "docx_scan.py"), CLEAN, "--layers", "ISP",
                  "--out", os.path.join(HERE, "scan-clean"), "--quiet")
    res2 = json.load(open(os.path.join(HERE, "scan-clean", "report.json"), encoding="utf-8"))
    left = [i for i in res2["issues"] if i["auto"]]
    check(not left, "清理后机械层零残留", "残留 %d 处" % len(left))

    print("\n── T6 保护区 ──")
    a, b = read_docx(SAMPLE), read_docx(CLEAN)
    changed = [n for n in a.order if not a.is_text_part(n) and a.data[n] != b.data[n]]
    check(not changed, "非文本部件逐字节未变", ",".join(changed) or "全部 %d 个部件" %
          len(a.protected_parts()))
    xml_changed = [n for n in a.order if a.is_text_part(n) and a.data[n] != b.data[n]]
    check(xml_changed, "文本部件确实变了", "、".join(xml_changed))
    # 只允许 <w:t> 的 body 变：把两边的 w:t 内容全挖掉后，XML 应完全一致
    def skeleton(buf: bytes) -> bytes:
        import re
        s = buf.decode("utf-8")
        return re.sub(r"(?<=>)[^<]*(?=</w:t>)", "", s).encode("utf-8")
    same = all(skeleton(a.data[n]) == skeleton(b.data[n]) for n in xml_changed)
    check(same, "XML 骨架未变（只有 <w:t> 里的字变了）")

    print("\n── T8 文本级保护区（掩码）──")
    from protect import masks_of as _masks

    def _covered(text):
        cov = set()
        for a, b in _masks(text):
            cov.update(range(a, b))
        return cov

    missed = []
    for text, token in MUST_MASK:
        cov = _covered(text)
        i = text.index(token)
        if not all(k in cov for k in range(i, i + len(token))):
            missed.append("%s（来自 %r）" % (token, text))
    check(not missed, "%d 个路径/命令/代码 token 整体受保护" % len(MUST_MASK),
          ("漏保护：%s" % "、".join(missed)) if missed else "一个字符都不许漏")

    over = [t for t in MUST_NOT_MASK if _masks(t)]
    check(not over, "%d 个该修的没被误保护" % len(MUST_NOT_MASK),
          ("被误保护：%s" % "、".join(over)) if over else "P01/P03/P04/P06/P07/P08/P09/S01/S02/S06 仍能命中")

    print("\n── T7 二次 apply 无事可做 ──")
    rc, out = run(os.path.join(ROOT, "scripts", "docx_apply.py"), CLEAN,
                  "--out", os.path.join(HERE, "clean2.docx"))
    check("无需改动" in out, "第二次 apply 报「无需改动」")

    if notes:
        print("\n── guard 明细 ──")
        print("\n".join(notes))

    print("\n" + ("=" * 56))
    if fails:
        print("❌ 自检失败 %d 项：" % len(fails))
        for f in fails:
            print("   · " + f)
        return 1
    print("✅ 自检全部通过（%d 条规则、%d 处命中、%d 处自动替换）"
          % (len(spec), len(res["issues"]), res["stats"]["auto"]))
    return 0


if __name__ == "__main__":
    sys.exit(main())
